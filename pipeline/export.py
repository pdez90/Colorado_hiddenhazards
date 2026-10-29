"""Write data/results/*: the sites workbook, GeoJSON and listing CSVs. The sheet/column layout is under main()."""
import pandas as pd, numpy as np, json
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
O='data/results/'
import re
status_label={'1 Cleanup / contamination list':'Cleanup / contamination list','2 TRI reporter':'TRI reporter','3 RCRA hazardous-waste handler':'RCRA hazardous-waste handler','4 Solid waste facility':'Solid waste facility','5 Other EPA permit only':'Other EPA permit only','6 Not on any list':'Not on any list'}
def main():
    s=pd.read_pickle('work/inv/sites_final.pkl'); df=pd.read_pickle('work/inv/records_final.pkl')
    # fill county by nearest zip centroid
    zc=pd.read_pickle('work/inv/zip_centroids.pkl'); zco=pd.read_csv('work/inv/zip_county.csv',dtype=str); zmap=dict(zip(zco.z,zco.COUNTY_NAME))
    from sklearn.neighbors import BallTree
    t=BallTree(np.radians(zc[['zlat','zlon']].values),metric='haversine')
    m=s.county.eq('')
    if m.any():
        _,ix=t.query(np.radians(s.loc[m,['lat','lon']].values),k=1); s.loc[m,'county']=[zmap.get(zc.zip5.iloc[i[0]],'') for i in ix]
    s['county']=s.county.str.replace(' COUNTY','',regex=False).str.title()
    s['radar']=s.radar_status.map(status_label)
    cols=['site_id','radar','match_basis','confidence','address','city','county','zip','lat','lon','geo_level','fe_sectors','ext_sectors','first_year','last_year','hazard_years','n_editions','n_establishments','n_hazard_establishments','earliest_est','hazard_firms','all_firms','industry_evidence',
          'cleanup_match','tri_match','rcra_handler_match','solid_waste_match','other_epa_match','cleanup_nearest','cleanup_nearest_m','tri_nearest','tri_nearest_m','rcra_handler_nearest','rcra_handler_nearest_m']
    for c in s.columns:
        if pd.api.types.is_object_dtype(s[c]) or pd.api.types.is_string_dtype(s[c]): s[c]=s[c].apply(lambda v: re.sub(r'^[=+@\-]+\s*','',v) if isinstance(v,str) else v)

    H=s[s.fe_hazard].sort_values(['radar','county','first_year'])[cols]
    E=s[s.ext_hazard&~s.fe_hazard].sort_values(['radar','county'])[cols]
    hidden=H[H.radar=='Not on any list'].sort_values(['confidence','county','first_year'],key=lambda c: c.map({'high':0,'medium':1,'low':2}) if c.name=='confidence' else c)
    cov=df.groupby('year').agg(listings=('name','size'),fe_hazard_listings=('fe_hazard','sum'),with_zip=('zip5',lambda z:(z.str.match(r'^8[01]\d{3}$')).mean()),located=('located','mean')).reset_index()
    vols=df.groupby('year').vol.agg(lambda v:', '.join(sorted(set(v)))).reset_index(); cov=cov.merge(vols,on='year')
    readme=[
     ['Colorado historical hazardous industrial sites, 1961–2002',''],
     ['Method','Frickel & Elliott (Sites Unseen, 2018) relic-site approach applied to the Colorado Directory of Manufacturers, 25 editions 1961–2002.'],
     ['Hazardous sectors (F&E)','SIC 28 chemicals, 29 petroleum refining, 30 rubber & plastics, 33 primary metals, 34 fabricated metals, 35 industrial machinery, 37 transportation equipment (NAICS 325, 324, 326, 331, 332, 333, 336 for 1999–2002).'],
     ['Extended sectors','Not in F&E but commonly contaminating: wood preserving 2491, leather tanning 3111, electronic components & circuit boards 3671–3679, batteries 3691–3692, photographic chemicals 3861, metal mining / ore milling 10xx.'],
     ['Site','All directory listings geocoded to street level within 30 m of each other, across all years.'],
     ['radar','Best match to a current inventory: Cleanup / contamination list (EPA SEMS/Superfund, ACRES brownfields, FRS state/federal cleanup programs; CDPHE VCUP, brownfields, Superfund, RCRA corrective action, institutional controls, UMTRA, federal facilities) — matched if within 75 m, inside the site boundary, or same name within 250 m or same address. TRI reporter (100 m or name/address). RCRA hazardous-waste handler (40 m or name/address). Other EPA permit only (40 m or name/address). Not on any list = not on the searched EPA and CDPHE HMWMD lists (FRS, TRI, SEMS/Superfund, RCRA, CDPHE cleanup, VCUP, brownfields, covenants); state petroleum-tank and release registers are shown separately on the Contamination & health tab of the app.'],
     ['match_basis','What the radar match rests on, strongest first: same entity (name or address) = an inventory record carries the hazardous firm\'s name (two shared informative words, or the whole name up to word order and OCR noise) or the site\'s street address; inside a cleanup boundary = the point lies in a mapped CDPHE cleanup-site polygon; nearby record only = a differently named record lies within the tier\'s radius (40–100 m) and nothing else links it to the site, so the record may belong to another occupant of the block (a pharmacy\'s RCRA notification, a car wash\'s permit); none = Not on any list.'],
     ['confidence','high = street-level geocode AND industry code printed in the directory; medium = one of the two; low = neither (inferred industry and approximate location).'],
     ['geo_level','address = number interpolated on the matching TIGER/Line 2023 street segment; address_approx = OCR-tolerant street match, neighbouring ZIP, or nearest segment end (verify before field work).'],
     ['Sources','Colorado Directory of Manufacturers (scans supplied by user); US EPA Facility Registry Service state file for Colorado (downloaded 16 Oct 2026); EPA TRI facility table (Envirofacts, 16 Oct 2026); CDPHE HMWMD layers, geodata.colorado.gov (16 Oct 2026); US Census TIGER/Line 2023 ADDRFEAT.'],
     ['Limits','OCR errors in names, numbers and ZIPs; 1961–1965 editions have no ZIP codes; 2002 covers only the second half of the alphabet of places (part 1 missing); 1971 and 1985 volumes appear partial; renumbered historical streets fall back to ZIP centroids (listed separately). A site absent from regulatory lists is a candidate for review, not evidence of contamination.'],
    ]
    with pd.ExcelWriter(O+'CO_hazardous_sites_1961-2002.xlsx',engine='openpyxl') as w:
        pd.DataFrame(readme,columns=['Item','Description']).to_excel(w,sheet_name='Read me',index=False)
        pd.DataFrame({'radar':['Cleanup / contamination list','TRI reporter','RCRA hazardous-waste handler','Solid waste facility','Other EPA permit only','Not on any list','Total']}).to_excel(w,sheet_name='Summary',index=False)
        hidden.to_excel(w,sheet_name='Hidden sites',index=False)
        H.to_excel(w,sheet_name='All F&E hazardous sites',index=False)
        E.to_excel(w,sheet_name='Extended-sector sites',index=False)
        cov.to_excel(w,sheet_name='Coverage by edition',index=False)
    wb=load_workbook(O+'CO_hazardous_sites_1961-2002.xlsx')
    ws=wb['Summary']
    hdr=['radar','F&E sites','share','high confidence','medium','low']
    for j,h in enumerate(hdr,1): ws.cell(1,j,h)
    n=len(H)+1
    for i in range(2,8):
        if True:
            ws.cell(i,2,f"=COUNTIF('All F&E hazardous sites'!$B$2:$B${n},A{i})")
            for j,c in zip((4,5,6),('high','medium','low')):
                ws.cell(i,j,f"=COUNTIFS('All F&E hazardous sites'!$B$2:$B${n},$A{i},'All F&E hazardous sites'!$D$2:$D${n},\"{c}\")")   # D = confidence (C = match_basis)
        ws.cell(i,3,f"=IFERROR(B{i}/$B$8,0)"); ws.cell(i,3).number_format='0.0%'
    ws.cell(8,2,"=SUM(B2:B7)"); 
    for j,c in zip((4,5,6),'DEF'): ws.cell(8,j,f"=SUM({c}2:{c}7)")
    ws.cell(8,3,"=IFERROR(B8/$B$8,0)"); ws.cell(8,3).number_format='0.0%'
    ws.cell(10,1,'Counts recalculate from the All F&E hazardous sites sheet.')
    for wsx in wb.worksheets:
        if wsx.title!='Summary':
            for row in wsx.iter_rows():
                for c in row:
                    if isinstance(c.value,str) and c.value[:1] in '=+@-': c.value=c.value.lstrip('=+@- ')
        for row in wsx.iter_rows():
            for c in row: c.font=Font(name='Arial',size=10,bold=(c.row==1))
        for c in wsx[1]: c.fill=PatternFill('solid',fgColor='DDE5DD')
        wsx.freeze_panes='A2'
        for j,col in enumerate(wsx.columns,1):
            L=max(len(str(c.value or ''))for c in list(col)[:200]); wsx.column_dimensions[get_column_letter(j)].width=min(max(10,L+2),60)
    wb['Read me'].column_dimensions['B'].width=120
    for r in wb['Read me'].iter_rows(min_row=2):
        r[1].alignment=Alignment(wrap_text=True,vertical='top')
    wb['Coverage by edition'].column_dimensions['F'].width=60
    for r in range(2,40):
        for c in ('D','E'): wb['Coverage by edition'][f'{c}{r}'].number_format='0%'
    wb.save(O+'CO_hazardous_sites_1961-2002.xlsx')
    # GIS + CSV
    gj={'type':'FeatureCollection','features':[{'type':'Feature','geometry':{'type':'Point','coordinates':[round(r.lon,6),round(r.lat,6)]},
       'properties':{k:(None if (isinstance(v,float) and np.isnan(v)) else (v.item() if hasattr(v,'item') else v)) for k,v in r.items() if k not in ('lat','lon')}} for _,r in s[cols[:-6]+['fe_hazard','ext_hazard']].iterrows()]}
    json.dump(gj,open(O+'CO_historical_industrial_sites.geojson','w'))
    R=df.copy(); R['sic']=R.sic.apply(lambda x:';'.join(x)); R['naics']=R.naics.apply(lambda x:';'.join(x) if isinstance(x,list) else '')
    rc=['year','vol','page','name','street','city_g','zip5','est','size','sic','sic_source','naics','fe_hazard','fe_sectors','ext_hazard','ext_sectors','estab_id','site_id','lat','lon','geo_level']
    R[rc].rename(columns={'city_g':'city','zip5':'zip','vol':'source_volume'}).to_csv(O+'CO_directory_listings_1961-2002.csv',index=False)
    U=R[R.fe_hazard & ~df.located][rc].rename(columns={'city_g':'city','zip5':'zip','vol':'source_volume'})
    U.to_csv(O+'CO_hazardous_listings_not_street_located.csv',index=False)
    s.to_pickle('work/inv/sites_export.pkl')
    print(len(H),len(hidden),len(E),len(R),len(U))

if __name__=='__main__': main()
